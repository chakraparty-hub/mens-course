# -*- coding: utf-8 -*-
"""Builds THE FILTER course talk-track document (.docx)."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import course_content as C

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GOLD = RGBColor(0xB8, 0x86, 0x0B)
DARK = RGBColor(0x26, 0x26, 0x26)
GREY = RGBColor(0x5A, 0x5A, 0x5A)
LIGHT = RGBColor(0xF4, 0xF1, 0xEA)

BODY_FONT = "Calibri"
HEAD_FONT = "Calibri"

doc = Document()

# ---------- base styles ----------
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.18

for hname, size, color in (("Heading 1", 17, NAVY), ("Heading 2", 13, NAVY), ("Heading 3", 11, GOLD)):
    st = doc.styles[hname]
    st.font.name = HEAD_FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(14 if hname != "Heading 1" else 18)
    st.paragraph_format.space_after = Pt(6)

def set_east_asia(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), HEAD_FONT)

def para(text="", style=None, align=None, space_before=None, space_after=None, size=None, color=None, bold=None, italic=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        set_east_asia(r)
    return p

def rich_para(segments, style=None, align=None, space_before=None, space_after=None):
    """segments: list of (text, kwargs) tuples."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = Pt(space_before)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    for text, kw in segments:
        r = p.add_run(text)
        if kw.get("size"):
            r.font.size = Pt(kw["size"])
        if kw.get("color"):
            r.font.color.rgb = kw["color"]
        if kw.get("bold"):
            r.bold = True
        if kw.get("italic"):
            r.italic = True
        if kw.get("font"):
            r.font.name = kw["font"]
        set_east_asia(r)
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.name = BODY_FONT
    if bold:
        r.bold = True
    if color:
        r.font.color.rgb = color
    set_east_asia(r)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)

def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- footer with page number ----------
section = doc.sections[0]
footer_p = section.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_p.add_run("THE FILTER — Coach Talk Track   |   Page ")
r.font.size = Pt(9)
r.font.color.rgb = GREY
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
r._element.addnext(fld)

# =========================================================
# TITLE PAGE
# =========================================================
para("", space_after=60)
p = para(C.COURSE_TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=44, color=NAVY, bold=True, space_after=6)
p = para(C.COURSE_SUBTITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=15, color=DARK, space_after=28)

tag = para(C.TAGLINE, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=GOLD, italic=True, space_after=40)

para("A complete coaching program with talk tracks — ready to teach live or record.",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GREY, space_after=4)
para("Includes: 8 modules · Talk tracks · Coaching exercises · Intake questions · Discovery-call script · Homework templates · 90-day roadmap · Marketing pack",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY, space_after=4)
para("For coaches of successful men who want the right woman — not the wrong one.",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY, space_after=4)

page_break()

# =========================================================
# HOW TO USE + OVERVIEW
# =========================================================
doc.add_heading("How to Use This Talk Track", level=1)
for item in C.HOW_TO_USE:
    para(item, style="List Bullet")

doc.add_heading("Course Overview", level=1)
for title, items in C.COURSE_OVERVIEW:
    rich_para([(title, {"bold": True, "color": NAVY, "size": 11})], space_before=6, space_after=4)
    for it in items:
        para(it, style="List Bullet")

# At-a-glance table
doc.add_heading("The 8 Modules at a Glance", level=2)
tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, htxt in enumerate(("Week", "Module", "Client outcome")):
    set_cell_text(hdr[i], htxt, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
    shade_cell(hdr[i], "1F3A5F")
for row in C.ROADMAP:
    cells = tbl.add_row().cells
    set_cell_text(cells[0], row[0], size=10)
    set_cell_text(cells[1], row[1], size=10, bold=True)
    set_cell_text(cells[2], row[2], size=10)

page_break()

# =========================================================
# MODULES
# =========================================================
for m in C.MODULES:
    doc.add_heading(f"Module {m['num']} — {m['title']}", level=1)

    rich_para([("Outcome: ", {"bold": True, "color": GOLD}), (m["outcome"], {})], space_after=8)

    doc.add_heading("Core Concepts", level=2)
    for title, desc in m["concepts"]:
        rich_para([(title + " — ", {"bold": True}), (desc, {})], style="List Bullet")

    doc.add_heading("Talk Track", level=2)
    for p_ in m["talk_track"]:
        para(p_, space_after=8)

    doc.add_heading("Coaching Exercise", level=2)
    rich_para([(m["exercise"]["title"], {"bold": True, "color": NAVY})], space_after=4)
    for i, step in enumerate(m["exercise"]["steps"], start=1):
        para(f"{i}. {step}", space_after=3)

    page_break()

# =========================================================
# COACHING TOOLKIT
# =========================================================
doc.add_heading("Coaching Toolkit", level=1)
para(C.TOOLKIT_INTRO, italic=True, color=GREY, space_after=10)

doc.add_heading("10 Intake Questions", level=2)
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
set_cell_text(hdr[0], "#", bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
set_cell_text(hdr[1], "Question", bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
shade_cell(hdr[0], "1F3A5F")
shade_cell(hdr[1], "1F3A5F")
for qid, q in C.INTAKE_QUESTIONS:
    cells = tbl.add_row().cells
    set_cell_text(cells[0], qid, size=10)
    set_cell_text(cells[1], q, size=10)

doc.add_heading("Discovery-Call Script", level=2)
for stage, text in C.DISCOVERY_CALL:
    rich_para([(stage + ": ", {"bold": True, "color": NAVY}), (text, {})], space_after=6)

doc.add_heading("Objection Handling", level=2)
for obj, reply in C.OBJECTIONS:
    rich_para([("Client: ", {"bold": True, "color": GOLD}), (obj, {})], space_after=2)
    rich_para([("You: ", {"bold": True, "color": GOLD}), (reply, {})], space_after=8)

doc.add_heading("Homework Templates (assign after each module)", level=2)
templates = [
    ("After Module 1", "The Magnet Audit — one-page list of what your success broadcasts and what it attracts."),
    ("After Module 2", "The Standard Worksheet — value inventory + five written non-negotiables."),
    ("After Module 3", "The Red Flag Log — 7 patterns + the rationalization you used for each."),
    ("After Module 4", "The Vetting Checklist — 90-day structure + your question bank."),
    ("After Module 5", "The Partnership Matrix — score her and yourself in five columns."),
    ("After Module 6", "Boundary Scripts — three statements rehearsed out loud."),
    ("After Module 7", "The Decision Sheet — head, heart, gut, with a decision date one week out."),
    ("After Module 8", "The Partnership Agreement + one-page mission-first statement."),
]
for when, what in templates:
    rich_para([(when + " — ", {"bold": True}), (what, {})], style="List Bullet")

doc.add_heading("90-Day Program Roadmap", level=2)
tbl = doc.add_table(rows=1, cols=3)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tbl.rows[0].cells
for i, htxt in enumerate(("Week", "Module", "Client outcome")):
    set_cell_text(hdr[i], htxt, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10)
    shade_cell(hdr[i], "1F3A5F")
for row in C.ROADMAP:
    cells = tbl.add_row().cells
    set_cell_text(cells[0], row[0], size=10)
    set_cell_text(cells[1], row[1], size=10, bold=True)
    set_cell_text(cells[2], row[2], size=10)

page_break()

# =========================================================
# MARKETING PACK
# =========================================================
doc.add_heading("Marketing Pack (Bonus)", level=1)

doc.add_heading("Elevator Pitch", level=2)
para(C.MARKETING["pitch"])

doc.add_heading("Pain-Point Bullets for Your Sales Page", level=2)
for b in C.MARKETING["pain_points"]:
    para(b, style="List Bullet")

doc.add_heading("Email Sequences", level=2)
for subj, body in C.MARKETING["emails"]:
    rich_para([("Subject: ", {"bold": True, "color": NAVY}), (subj, {"bold": True})], space_after=2)
    para(body, space_after=8)

doc.add_heading("Social Media Captions", level=2)
for cap in C.MARKETING["captions"]:
    para(cap, style="List Bullet")

# ---------- closing line ----------
para("", space_before=16)
para("End of course. Teach it, record it, make it yours.",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GOLD, italic=True)

OUT = "The_Filter_Coach_Talk_Track.docx"
doc.save(OUT)
print("saved:", OUT)
